from collections import defaultdict
import logging
import csv
import io

from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, AllowAny
from rest_framework import status
from rest_framework.pagination import PageNumberPagination
from django.db import transaction
from django.db.models import Count, Q
from django.core.exceptions import ObjectDoesNotExist
from django.utils import timezone
from datetime import datetime, time, date
import calendar
from openpyxl import load_workbook
from rest_framework.parsers import MultiPartParser, FormParser

# --- Models ---
from .models import Attendance, StaffAttendance
from teachers.models import Teacher
from students.models import Student, Enrollment
from school.models import AcademicYear
from academics.models import Section, ClassTeacher, Standard  # <--- IMPORTED ClassTeacher
from staff.models import NonTeachingStaff
from .permissions import IsAdmin, IsTeacher

# --- Mixins ---
from school.mixins import EnrollmentYearAccessMixin

# --- Serializers ---
from .serializers import (
    StudentAttendanceHistorySerializer, 
    DailyAttendanceListSerializer,
    StaffMarkAttendanceSerializer, 
    StaffAttendanceHistorySerializer
)
from staff.permissions import IsStaff, IsOwnerAdminOrAdminStaff
from notifications.whatsapp import send_whatsapp_template


logger = logging.getLogger(__name__)


WHATSAPP_DISABLED_PREFIX = 'WhatsApp alert disabled:'


def format_attendance_alert_date(date_value):
    if hasattr(date_value, 'strftime'):
        return date_value.strftime('%d %b %Y')

    try:
        parsed_date = datetime.strptime(str(date_value), '%Y-%m-%d').date()
        return parsed_date.strftime('%d %b %Y')
    except (TypeError, ValueError):
        return str(date_value)


def send_student_attendance_whatsapp(enrollment, attendance_status, attendance_date):
    student = enrollment.student
    parent_phone = getattr(student, 'father_phone', '') or getattr(student, 'mother_phone', '')
    if not parent_phone:
        return False, {'error': 'Parent phone number is missing.'}

    section = enrollment.section
    class_label = enrollment.standard.name
    if section:
        class_label = f'{class_label}-{section.name}'

    message = (
        f'Attendance update: {student.student_name} ({student.student_id}) '
        f'is marked {attendance_status} for {format_attendance_alert_date(attendance_date)}. '
        f'Class: {class_label}.'
    )
    return send_whatsapp_template(parent_phone, message, alert_type='attendance')


def normalize_import_key(value):
    return str(value or '').strip().lower().replace(' ', '_').replace('-', '_')


def normalize_import_status(value):
    raw = str(value or '').strip().lower().replace('_', ' ').replace('-', ' ')
    compact = raw.replace(' ', '')
    if compact in {'present', 'p'}:
        return 'Present'
    if compact in {'absent', 'a'}:
        return 'Absent'
    if compact in {'late', 'l'}:
        return 'Late'
    if compact in {'onleave', 'leave'}:
        return 'On Leave'
    raise ValueError("Status must be Present, Absent, Late, or On Leave.")


def parse_import_date(value):
    if isinstance(value, date):
        return value
    if isinstance(value, datetime):
        return value.date()

    raw = str(value or '').strip()
    if not raw:
        raise ValueError('Date is required.')

    for fmt in (
        '%Y-%m-%d',
        '%d-%m-%Y',
        '%d/%m/%Y',
        '%m/%d/%Y',
        '%d.%m.%Y',
        '%d %b %Y',
        '%d %B %Y',
        '%Y/%m/%d',
    ):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Invalid date format '{raw}'.")


def read_uploaded_rows(uploaded_file):
    filename = (uploaded_file.name or '').lower()
    if filename.endswith('.csv'):
        raw = uploaded_file.read().decode('utf-8-sig')
        reader = csv.reader(io.StringIO(raw))
        rows = [row for row in reader if any(str(cell).strip() for cell in row)]
        if not rows:
            raise ValueError('CSV file is empty.')
        headers = [normalize_import_key(cell) for cell in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]]

    if filename.endswith('.xlsx') or filename.endswith('.xlsm'):
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        sheet = workbook.active
        rows = [
            [cell for cell in row]
            for row in sheet.iter_rows(values_only=True)
            if any(str(cell).strip() for cell in row if cell is not None)
        ]
        if not rows:
            raise ValueError('Excel file is empty.')
        headers = [normalize_import_key(cell) for cell in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]]

    raise ValueError('Only CSV and XLSX files are supported.')


def first_existing_value(row, *keys):
    for key in keys:
        value = row.get(normalize_import_key(key))
        if value is not None and str(value).strip():
            return str(value).strip()
    return ''


# ==========================================
# 1. MARK ATTENDANCE (Class Teacher Only)
# ==========================================
class TeacherAttendanceView(APIView, EnrollmentYearAccessMixin):
    permission_classes = [IsAuthenticated]

    def get_teacher_section(self, user):
        """
        Helper: Finds the section assigned to the Class Teacher for the ACTIVE Academic Year.
        """
        try:
            teacher = Teacher.objects.get(user=user)
            
            # Get Active Year
            active_year = self.get_active_year()
            if not active_year: 
                return None
            
            # Lookup in ClassTeacher Table
            class_teacher_record = ClassTeacher.objects.get(
                teacher=teacher, 
                academic_year=active_year
            )
            return class_teacher_record.section

        except (Teacher.DoesNotExist, ClassTeacher.DoesNotExist):
            return None

    def get(self, request):
        """FETCH ROLL CALL LIST"""
        section = self.get_teacher_section(request.user)
        if not section:
            return Response({"error": "You are not assigned as a Class Teacher for the active year."}, status=403)

        active_year = self.get_active_year()
        if not active_year:
            return Response({"error": "No active academic year configured."}, status=500)

        enrollments = Enrollment.objects.filter(
            section=section, 
            academic_year=active_year,
            is_active=True
        ).select_related('student').order_by('student__student_name')
        
        student_data = [
            {
                "student_id": e.student.student_id, 
                "name": e.student.student_name, 
                "roll_no": getattr(e.student, 'roll_no', idx+1)
            } 
            for idx, e in enumerate(enrollments)
        ]
        
        return Response({
            "class": section.standard.name,
            "section": section.name,
            "total_students": len(enrollments),
            "students": student_data
        }, status=200)

    def post(self, request):
        """SUBMIT BATCH ATTENDANCE"""
        section = self.get_teacher_section(request.user)
        if not section:
            return Response({"error": "You are not assigned as a Class Teacher for the active year."}, status=403)

        date_val = request.data.get("date")
        attendance_list = request.data.get("attendance_list") 

        if not date_val or not attendance_list:
            return Response({"error": "Date and attendance_list are required"}, status=400)

        active_year = self.get_active_year()
        if not active_year:
             return Response({"error": "No Active Academic Year found."}, status=500)

        try:
            attendance_alerts = []
            with transaction.atomic():
                # Clear old records for this date/section/year to prevent duplicates
                Attendance.objects.filter(
                    enrollment__section=section, 
                    date=date_val,
                    enrollment__academic_year=active_year 
                ).delete()

                records = []
                for item in attendance_list:
                    sid = item.get("student_id")
                    status_val = item.get("status") 
                    
                    enrollment = Enrollment.objects.filter(
                        student__student_id=sid, 
                        section=section,
                        academic_year=active_year
                    ).first()

                    if enrollment:
                        records.append(Attendance(
                            enrollment=enrollment,
                            date=date_val,
                            status=status_val
                        ))
                        attendance_alerts.append((enrollment, status_val))
                
                Attendance.objects.bulk_create(records)

            whatsapp_sent = 0
            whatsapp_failed = 0
            whatsapp_skipped = 0
            for enrollment, attendance_status in attendance_alerts:
                ok, response_data = send_student_attendance_whatsapp(
                    enrollment,
                    attendance_status,
                    date_val,
                )
                if ok:
                    whatsapp_sent += 1
                elif str(response_data.get('error', '')).startswith(WHATSAPP_DISABLED_PREFIX):
                    whatsapp_skipped += 1
                else:
                    whatsapp_failed += 1
                    logger.warning(
                        'Student attendance WhatsApp failed for %s: %s',
                        enrollment.student.student_id,
                        response_data,
                    )

            return Response({
                "message": "Attendance marked successfully!",
                "whatsapp_sent": whatsapp_sent,
                "whatsapp_failed": whatsapp_failed,
                "whatsapp_skipped": whatsapp_skipped,
            }, status=200)
        except Exception as e:
            return Response({"error": str(e)}, status=400)


# ==========================================
# 2. STUDENT HISTORY (Calendar UI)
# ==========================================
class StudentAttendanceHistoryView(APIView, EnrollmentYearAccessMixin):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        target_student_id = request.query_params.get('student_id')
        year_param = request.query_params.get('year') # Expecting "2023-2024"

        if not target_student_id or not year_param:
            return Response({"error": "student_id and year are required"}, status=400)

        user = request.user

        # Security Checks
        if user.user_type == 'student' and user.username != target_student_id:
            return Response({"error": "Access Denied"}, status=403)
        
        # 1. Get Allowed Years
        allowed_years = self.get_allowed_academic_years(user)

        # 2. Teacher Access Control
        if user.user_type == 'teacher':
            try:
                teacher = Teacher.objects.get(user=user)
                active_year = self.get_active_year()
                
                # Check ClassTeacher allocation
                ct_record = ClassTeacher.objects.get(teacher=teacher, academic_year=active_year)
                my_section = ct_record.section

                # Check if student is enrolled in that section for valid years
                is_student_in_class = Enrollment.objects.filter(
                    student__student_id=target_student_id,
                    section=my_section,
                    academic_year__in=allowed_years
                ).exists()

                if not is_student_in_class:
                      return Response({"error": "Student not in your class (Active Year Only)."}, status=403)
            except ClassTeacher.DoesNotExist:
                return Response({"error": "You are not a Class Teacher for the active year."}, status=403)
            except Exception:
                return Response({"error": "Teacher/Student error"}, status=403)

        # 3. FETCH DATA
        records = Attendance.objects.filter(
            enrollment__student__student_id=target_student_id,
            enrollment__academic_year__in=allowed_years,  # Security Lock
            enrollment__academic_year__name=year_param    # User Filter
        ).order_by('date')

        student_name = (
            Student.objects.filter(student_id=target_student_id)
            .values_list('student_name', flat=True)
            .first()
        )

        enrollment_for_year = Enrollment.objects.filter(
            student__student_id=target_student_id,
            academic_year__name=year_param
        ).select_related('standard', 'section').first()

        # --- GROUPING LOGIC ---
        grouped_data = {}
        annual_stats = {"present": 0, "absent": 0, "late": 0}

        for record in records:
            month_key = str(record.date.month)
            status_key = record.status.lower()

            if month_key not in grouped_data:
                grouped_data[month_key] = {
                    "stats": {"present": 0, "absent": 0, "late": 0},
                    "dates": []
                }

            if status_key in grouped_data[month_key]["stats"]:
                grouped_data[month_key]["stats"][status_key] += 1
            if status_key in annual_stats:
                annual_stats[status_key] += 1

            grouped_data[month_key]["dates"].append({
                "date": record.date.day,
                "status": record.status,
                "status_display": record.get_status_display()
            })

        # Keep percentage defined even when there are no attendance records.
        total_days = annual_stats["present"] + annual_stats["absent"] + annual_stats["late"]
        percentage = (annual_stats["present"] / total_days * 100) if total_days > 0 else 0.0

        return Response({
            "student_id": target_student_id,
            "student_name": student_name or "",
            "class": enrollment_for_year.standard.name if enrollment_for_year and enrollment_for_year.standard else "",
            "section": enrollment_for_year.section.name if enrollment_for_year and enrollment_for_year.section else "",
            "year": year_param, 
            "attendance_percentage": f"{round(percentage, 1)}%",
            "annual_summary": annual_stats,
            "calendar_data": grouped_data
        }, status=200)


# ==========================================
# 3. CLASS DAILY SUMMARY
# ==========================================
class ClassDailySummaryView(APIView, EnrollmentYearAccessMixin):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        class_name = request.query_params.get('class')
        section_name = request.query_params.get('section')
        date_str = request.query_params.get('date')

        if not all([class_name, section_name, date_str]):
             return Response({"error": "class, section, and date are required"}, status=400)
        
        allowed_years = self.get_allowed_academic_years(request.user)

        # Teacher Permission Check
        if request.user.user_type == 'teacher':
            try:
                teacher = Teacher.objects.get(user=request.user)
                active_year = self.get_active_year()
                
                ct_record = ClassTeacher.objects.get(teacher=teacher, academic_year=active_year)
                my_section = ct_record.section

                if not (my_section.standard.name == class_name and 
                        my_section.name == section_name):
                    return Response({"error": "Access Denied"}, status=403)
            except ClassTeacher.DoesNotExist:
                return Response({"error": "You are not a Class Teacher for the active year."}, status=403)
            except:
                return Response({"error": "Profile error"}, status=403)

        records = Attendance.objects.filter(
            enrollment__standard__name=class_name,
            enrollment__section__name=section_name,
            enrollment__academic_year__in=allowed_years,
            date=date_str
        )

        stats = records.aggregate(
            present=Count('id', filter=Q(status__iexact='present')),
            absent=Count('id', filter=Q(status__iexact='absent')),
            late=Count('id', filter=Q(status__iexact='late')),
            on_leave=Count('id', filter=Q(status__iexact='on leave')),
        )

        serializer = DailyAttendanceListSerializer(records, many=True)

        return Response({
            "date": date_str,
            "class": f"{class_name}-{section_name}",
            "summary": stats,
            "students": serializer.data
        })


# ==========================================
# 4. UPDATE ATTENDANCE (Single Edit)
# ==========================================
class UpdateAttendanceView(APIView, EnrollmentYearAccessMixin):
    permission_classes = [IsAuthenticated]

    def put(self, request):
        user = request.user
        data = request.data
        
        sid = data.get('student_id')
        date_str = data.get('date')
        new_status = data.get('status')

        if not all([sid, date_str, new_status]):
             return Response({"error": "Missing fields: student_id, date, status"}, status=400)

        if user.user_type == 'student':
             return Response({"error": "Unauthorized"}, status=403)
        
        allowed_years = self.get_allowed_academic_years(user)

        try:
            enrollment = Enrollment.objects.get(
                student__student_id=sid, 
                academic_year__in=allowed_years
            )

            if user.user_type == 'teacher':
                try:
                    teacher = Teacher.objects.get(user=user)
                    active_year = self.get_active_year()

                    # Find assigned section via ClassTeacher
                    ct_record = ClassTeacher.objects.get(teacher=teacher, academic_year=active_year)
                    my_section = ct_record.section
                    
                    if enrollment.section != my_section:
                        return Response({"error": "Student not in your class."}, status=403)
                
                except Teacher.DoesNotExist:
                     return Response({"error": "Teacher profile not found"}, status=400)
                except ClassTeacher.DoesNotExist:
                     return Response({"error": "You are not a Class Teacher for the active year."}, status=403)

            record, created = Attendance.objects.update_or_create(
                enrollment=enrollment,
                date=date_str,
                defaults={'status': new_status}
            )

            whatsapp_ok, whatsapp_response = send_student_attendance_whatsapp(
                enrollment,
                new_status,
                date_str,
            )
            whatsapp_disabled = str(whatsapp_response.get('error', '')).startswith(
                WHATSAPP_DISABLED_PREFIX
            )
            if not whatsapp_ok and not whatsapp_disabled:
                logger.warning(
                    'Student attendance WhatsApp failed for %s: %s',
                    enrollment.student.student_id,
                    whatsapp_response,
                )

            action = "Created" if created else "Updated"
            return Response({
                "message": f"Attendance {action} to {new_status}",
                "whatsapp_sent": 1 if whatsapp_ok else 0,
                "whatsapp_failed": 0 if whatsapp_ok or whatsapp_disabled else 1,
                "whatsapp_skipped": 1 if whatsapp_disabled else 0,
            }, status=200)

        except Enrollment.DoesNotExist:
             return Response({"error": "Student not enrolled in an editable academic year."}, status=404)
        except Exception as e:
             return Response({"error": str(e)}, status=500)


class StudentWeeklyAttendanceImportView(APIView, EnrollmentYearAccessMixin):
    permission_classes = [IsAuthenticated, IsTeacher]
    parser_classes = (MultiPartParser, FormParser)

    def post(self, request):
        uploaded_file = request.FILES.get('file')
        if not uploaded_file:
            return Response({"error": "file is required"}, status=400)

        try:
            teacher = Teacher.objects.get(user=request.user)
            active_year = self.get_active_year()
            if not active_year:
                return Response({"error": "No Active Academic Year found."}, status=500)

            ct_record = ClassTeacher.objects.get(teacher=teacher, academic_year=active_year)
            rows = read_uploaded_rows(uploaded_file)
        except Teacher.DoesNotExist:
            return Response({"error": "Teacher profile not found."}, status=404)
        except ClassTeacher.DoesNotExist:
            return Response({"error": "You are not assigned as a Class Teacher for the active year."}, status=403)
        except ValueError as e:
            return Response({"error": str(e)}, status=400)

        if not rows:
            return Response({"error": "No attendance rows found in the file."}, status=400)

        created = 0
        updated = 0
        skipped = 0
        errors = []

        with transaction.atomic():
            for index, row in enumerate(rows, start=2):
                try:
                    student_id = first_existing_value(row, 'student_id', 'id')
                    attendance_date = parse_import_date(row.get('date'))
                    status_value = normalize_import_status(row.get('status'))

                    if not student_id:
                        raise ValueError('student_id is required')

                    enrollment = Enrollment.objects.filter(
                        student__student_id=student_id,
                        section=ct_record.section,
                        academic_year=active_year,
                    ).first()

                    if not enrollment:
                        skipped += 1
                        errors.append(f'Row {index}: student {student_id} not found in your class.')
                        continue

                    _, was_created = Attendance.objects.update_or_create(
                        enrollment=enrollment,
                        date=attendance_date,
                        defaults={'status': status_value},
                    )
                    if was_created:
                        created += 1
                    else:
                        updated += 1
                except Exception as e:
                    skipped += 1
                    errors.append(f'Row {index}: {e}')

        return Response({
            "status": 200,
            "message": "Student attendance imported successfully.",
            "summary": {
                "created": created,
                "updated": updated,
                "skipped": skipped,
                "total_rows": len(rows),
            },
            "errors": errors[:20],
        }, status=200)


class TeacherAttendanceBulkImportView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]
    parser_classes = (MultiPartParser, FormParser)

    def post(self, request):
        uploaded_file = request.FILES.get('file')
        if not uploaded_file:
            return Response({"error": "file is required"}, status=400)

        try:
            rows = read_uploaded_rows(uploaded_file)
        except ValueError as e:
            return Response({"error": str(e)}, status=400)

        if not rows:
            return Response({"error": "No attendance rows found in the file."}, status=400)

        created = 0
        updated = 0
        skipped = 0
        errors = []

        with transaction.atomic():
            for index, row in enumerate(rows, start=2):
                try:
                    teacher_id = first_existing_value(row, 'teacher_id', 'id')
                    attendance_date = parse_import_date(row.get('date'))
                    status_value = normalize_import_status(row.get('status'))

                    if not teacher_id:
                        raise ValueError('teacher_id is required')

                    teacher = Teacher.objects.filter(teacher_id=teacher_id).first()
                    if not teacher:
                        skipped += 1
                        errors.append(f'Row {index}: teacher {teacher_id} not found.')
                        continue

                    _, was_created = TeacherAttendance.objects.update_or_create(
                        teacher=teacher,
                        date=attendance_date,
                        defaults={'status': status_value},
                    )
                    if was_created:
                        created += 1
                    else:
                        updated += 1
                except Exception as e:
                    skipped += 1
                    errors.append(f'Row {index}: {e}')

        return Response({
            "status": 200,
            "message": "Teacher attendance imported successfully.",
            "summary": {
                "created": created,
                "updated": updated,
                "skipped": skipped,
                "total_rows": len(rows),
            },
            "errors": errors[:20],
        }, status=200)


# ==========================================
# 5. TEACHER VIEW CLASS ATTENDANCE (Register)
# ==========================================
class TeacherViewClassAttendance(APIView, EnrollmentYearAccessMixin):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        date_str = request.query_params.get('date')

        if not date_str:
            return Response({"error": "Date parameter is required"}, status=400)

        if user.user_type != 'teacher':
             return Response({"error": "Unauthorized. Teachers only."}, status=403)
        
        try:
            teacher_profile = Teacher.objects.get(user=user)
            active_year = self.get_active_year()
            if not active_year:
                return Response({"error": "No Active Academic Year found"}, status=500)

            # Get Section via ClassTeacher
            ct_record = ClassTeacher.objects.get(teacher=teacher_profile, academic_year=active_year)
            my_section = ct_record.section

        except Teacher.DoesNotExist:
             return Response({"error": "Teacher profile not found"}, status=404)
        except ClassTeacher.DoesNotExist:
             return Response({"error": "You are not assigned as a Class Teacher for the current year."}, status=403)

        allowed_years = self.get_allowed_academic_years(user)

        enrollments = Enrollment.objects.filter(
            section=my_section,
            academic_year__in=allowed_years
        ).select_related('student')
        
        attendance_records = Attendance.objects.filter(
            enrollment__section=my_section,
            enrollment__academic_year__in=allowed_years, 
            date=date_str
        )
        
        attendance_map = {record.enrollment.id: record.status for record in attendance_records}
        
        stats = {"Present": 0, "Absent": 0, "Late": 0, "On Leave": 0, "Not Marked": 0}

        final_list = []
        for idx, enroll in enumerate(enrollments):
            status_val = attendance_map.get(enroll.id, "Not Marked")
            
            stats_key = status_val.title() if status_val.lower() in ["present", "absent", "late", "on leave"] else "Not Marked"
            if stats_key in stats:
                stats[stats_key] += 1
            
            final_list.append({
                "student_id": enroll.student.student_id,
                "name": enroll.student.student_name,
                "roll_no": getattr(enroll.student, 'roll_no', idx+1),
                "status": status_val
            })

        return Response({
            "date": date_str,
            "class": my_section.standard.name,
            "section": my_section.name,
            "summary": stats,
            "attendance_data": final_list
        }, status=200)


# ==========================================
# 6. STUDENT SELF VIEW
# ==========================================
class StudentMyAttendanceView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            student = request.user.student_profile
        except AttributeError:
            return Response({"error": "Access Denied. Students only."}, 403)
        
        try:
            active_year = AcademicYear.objects.get(is_current=True)
        except AcademicYear.DoesNotExist:
             return Response({"error": "System Configuration Error: No Active Year."}, status=500)

        base_query = Attendance.objects.filter(
            enrollment__student=student, 
            enrollment__academic_year=active_year 
        )

        req_date = request.query_params.get('date')
        if req_date:
            try:
                record = base_query.get(date=req_date)
                return Response({
                    "date": req_date,
                    "status": record.status,
                    "status_display": record.get_status_display()
                })
            except Attendance.DoesNotExist:
                return Response({"date": req_date, "status": "Not Marked", "status_display": "-"})

        req_month = request.query_params.get('month')
        if req_month:
            records = base_query.filter(
                date__month=req_month
            ).order_by('date')

            monthly_stats = {"Present": 0, "Absent": 0, "Late": 0, "On Leave": 0}
            calendar_data = []

            for r in records:
                status_key = r.status.title() if r.status.lower() in ["present", "absent", "late", "on leave"] else "Other"
                if status_key in monthly_stats: monthly_stats[status_key] += 1
                
                calendar_data.append({
                    "date": r.date,
                    "day": r.date.strftime("%A"),
                    "status": r.status
                })
            
            return Response({
                "view": "Monthly",
                "month": req_month,
                "year": active_year.name,
                "total_working_days": records.count(),
                "summary": monthly_stats,
                "history": calendar_data
            })

        records = base_query.order_by('date')
        grouped_data = {}
        annual_summary = {"present": 0, "absent": 0, "late": 0, "on leave": 0}
        total_present_count = 0
        total_working_days = records.count()

        for record in records:
            status_key = record.status.lower() 
            month_key = str(record.date.month)
            if status_key in annual_summary: annual_summary[status_key] += 1
            if status_key == 'present': total_present_count += 1
            
            if month_key not in grouped_data:
                grouped_data[month_key] = {"stats": {"present": 0, "absent": 0, "late": 0, "on leave": 0}, "dates": []}
            if status_key in grouped_data[month_key]["stats"]: grouped_data[month_key]["stats"][status_key] += 1
            
            grouped_data[month_key]["dates"].append({
                "date": record.date.day,
                "status": status_key,
                "status_display": record.get_status_display().lower()
            })

        percentage = (total_present_count / total_working_days * 100) if total_working_days > 0 else 0.0

        return Response({
            "student_id": student.student_id,
            "year": active_year.name,
            "academic_year": {
                "id": active_year.id,
                "name": active_year.name,
                "start_date": active_year.start_date,
                "end_date": active_year.end_date,
            },
            "attendance_percentage": f"{round(percentage, 1)}%",
            "annual_summary": annual_summary,
            "calendar_data": grouped_data
        })

##### teachers and staff attendance starts from here 

from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from schooladmin.permissions import IsAdmin
from django.utils import timezone
from django.shortcuts import get_object_or_404
from django.db.models import Q
from django.core import signing
from math import radians, cos, sin, asin, sqrt
from datetime import date, timedelta
import calendar

from .models import StaffAttendance, TeacherAttendance, AttendanceConfig, QRAttendanceSession, QRAttendanceScanLog
from .serializers import (
    AttendanceConfigSerializer,
    StaffAttendanceHistorySerializer,
    TeacherAttendanceHistorySerializer,
    DailyStaffStatusSerializer,
    DailyTeacherStatusSerializer,
    QRAttendanceSessionStartSerializer,
    QRAttendanceSessionSerializer,
    QRAttendanceScanSerializer,
)
from staff.models import NonTeachingStaff
from teachers.models import Teacher
from holidays.models import Holiday

# ==========================================
# HELPER: GEOFENCE CHECK
# ==========================================
def check_geofence_and_time(user_lat, user_long):
    config = AttendanceConfig.objects.first()
    if not config: return False, "Config Missing"

    try:
        R = 6371000
        lat1, lon1 = radians(float(user_lat)), radians(float(user_long))
        lat2, lon2 = radians(float(config.school_latitude)), radians(float(config.school_longitude))
        dphi, dlambda = lat2 - lat1, lon2 - lon1
        a = sin(dphi/2)**2 + cos(lat1)*cos(lat2)*sin(dlambda/2)**2
        dist = R * (2 * asin(sqrt(a)))

        if dist > config.allowed_radius_meters:
            return False, f"Out of range ({int(dist)}m)"
    except:
        return False, "Invalid Coords"

    status = "Late" if timezone.now().time() > config.late_cutoff_time else "Present"
    return True, status


# ==========================================
# VIEW 1: STAFF OWN HISTORY
# ==========================================
class StaffAttendanceHistoryView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        staff_profile = request.user.staff_profile
        param_date = request.query_params.get('date')
        param_month = request.query_params.get('month')
        param_year = request.query_params.get('year')

        today = date.today()
        year = int(param_year) if param_year else today.year

        if param_date:
            try:
                record = StaffAttendance.objects.get(staff=staff_profile, date=param_date)
                data = StaffAttendanceHistorySerializer(record).data
                return Response({"period": "Single Date", "data": data}, 200)
            except StaffAttendance.DoesNotExist:
                return Response({"period": "Single Date", "date": param_date, "status": "Absent"}, 200)

        if param_year and not param_month:
            start_date, end_date = date(year, 1, 1), (today if year == today.year else date(year, 12, 31))
            period_label = f"Year {year}"
        else:
            month = int(param_month) if param_month else today.month
            _, days_in_month = calendar.monthrange(year, month)
            start_date = date(year, month, 1)
            limit_day = today.day if (year == today.year and month == today.month) else days_in_month
            end_date = date(year, month, limit_day)
            period_label = f"{calendar.month_name[month]} {year}"

        records = StaffAttendance.objects.filter(staff=staff_profile, date__range=[start_date, end_date])
        attendance_map = {r.date: r for r in records}
        holiday_objs = Holiday.objects.filter(date__range=[start_date, end_date]).filter(Q(applicable_for='everyone') | Q(applicable_for='staff'))
        holiday_list = list(holiday_objs.values_list('date', flat=True))

        grouped_history, s_cnt, h_cnt, p_cnt, l_cnt, a_cnt = {}, 0, 0, 0, 0, 0
        curr = start_date
        while curr <= end_date:
            m_key = str(curr.month)
            if m_key not in grouped_history: grouped_history[m_key] = []
            
            rec = attendance_map.get(curr)
            if rec:
                status, check_in = rec.status, rec.check_in_time.strftime("%H:%M:%S")
                if status == 'Present': p_cnt += 1
                else: l_cnt += 1
            elif curr.weekday() == 6: status, check_in, s_cnt = "Sunday", None, s_cnt + 1
            elif curr in holiday_list: status, check_in, h_cnt = "Holiday", None, h_cnt + 1
            else: status, check_in, a_cnt = "Absent", None, a_cnt + 1

            grouped_history[m_key].append({"date": curr.strftime("%Y-%m-%d"), "check_in_time": check_in, "status": status})
            curr += timedelta(days=1)

        total_days = (end_date - start_date).days + 1
        w_days = total_days - s_cnt - h_cnt
        return Response({
            "period": period_label,
            "summary": {
                "total_days_passed": total_days, "sundays": s_cnt, "holidays": h_cnt, 
                "holiday_dates": [h.strftime("%Y-%m-%d") for h in holiday_list],
                "actual_working_days": max(0, w_days), "present": p_cnt, "late": l_cnt, "absent": a_cnt,
                "percentage": f"{round(((p_cnt+l_cnt)/w_days)*100,1)}%" if w_days > 0 else "0%"
            },
            "history": grouped_history
        }, 200)

# ==========================================
# VIEW 2: TEACHER OWN HISTORY
# ==========================================
class TeacherAttendanceHistoryView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        teacher_profile = request.user.teacher_profile
        param_date, param_month, param_year = request.query_params.get('date'), request.query_params.get('month'), request.query_params.get('year')
        today = date.today()
        year = int(param_year) if param_year else today.year
        joining_date = teacher_profile.joining_date
        teacher_profile_data = {
            "teacher_id": teacher_profile.teacher_id,
            "teacher_name": teacher_profile.name,
            "department": teacher_profile.department,
            "joining_date": joining_date.isoformat() if joining_date else None,
        }

        if param_date:
            try:
                target_date = datetime.strptime(param_date, "%Y-%m-%d").date()
            except ValueError:
                return Response({"error": "Invalid date format. Use YYYY-MM-DD."}, 400)

            if joining_date and target_date < joining_date:
                return Response({
                    "period": "Single Date",
                    "teacher_profile": teacher_profile_data,
                    "date": param_date,
                    "status": "Not Joined",
                    "message": f"Attendance starts from joining date {joining_date.isoformat()}."
                }, 200)

            try:
                record = TeacherAttendance.objects.get(teacher=teacher_profile, date=param_date)
                return Response({
                    "period": "Single Date",
                    "teacher_profile": teacher_profile_data,
                    "data": TeacherAttendanceHistorySerializer(record).data
                }, 200)
            except TeacherAttendance.DoesNotExist:
                if target_date.weekday() == 6:
                    return Response({"period": "Single Date", "teacher_profile": teacher_profile_data, "date": param_date, "status": "Sunday"}, 200)

                is_holiday = Holiday.objects.filter(date=target_date).filter(
                    Q(applicable_for='everyone') | Q(applicable_for='teachers')
                ).exists()
                if is_holiday:
                    return Response({"period": "Single Date", "teacher_profile": teacher_profile_data, "date": param_date, "status": "Holiday"}, 200)

                return Response({"period": "Single Date", "teacher_profile": teacher_profile_data, "date": param_date, "status": "Absent"}, 200)

        if param_year and not param_month:
            start_date, end_date = date(year, 1, 1), (today if year == today.year else date(year, 12, 31))
            period_label = f"Year {year}"
        else:
            month = int(param_month) if param_month else today.month
            _, dim = calendar.monthrange(year, month)
            start_date, end_date = date(year, month, 1), date(year, month, today.day if (year==today.year and month==today.month) else dim)
            period_label = f"{calendar.month_name[month]} {year}"

        requested_start_date = start_date
        if joining_date and start_date < joining_date:
            start_date = joining_date

        if start_date > end_date:
            return Response({
                "period": period_label,
                "teacher_profile": teacher_profile_data,
                "summary": {
                    "total_days_passed": 0,
                    "sundays": 0,
                    "holidays": 0,
                    "holiday_dates": [],
                    "actual_working_days": 0,
                    "present": 0,
                    "late": 0,
                    "absent": 0,
                    "percentage": "0%",
                    "attendance_percentage": "0%",
                    "joining_date": joining_date.isoformat() if joining_date else None,
                    "calculated_from": joining_date.isoformat() if joining_date else requested_start_date.isoformat(),
                    "message": "Attendance is calculated from the teacher joining date."
                },
                "history": {}
            }, 200)

        records = TeacherAttendance.objects.filter(teacher=teacher_profile, date__range=[start_date, end_date])
        att_map = {r.date: r for r in records}
        hols = Holiday.objects.filter(date__range=[start_date, end_date]).filter(Q(applicable_for='everyone') | Q(applicable_for='teachers'))
        hol_list = list(hols.values_list('date', flat=True))

        history, s_cnt, h_cnt, p_cnt, l_cnt, a_cnt = {}, 0, 0, 0, 0, 0
        curr = start_date
        while curr <= end_date:
            m_key = str(curr.month)
            if m_key not in history: history[m_key] = []
            rec = att_map.get(curr)
            if rec:
                status, cin = rec.status, rec.check_in_time.strftime("%H:%M:%S")
                if status == 'Present': p_cnt += 1
                else: l_cnt += 1
            elif curr.weekday() == 6: status, cin, s_cnt = "Sunday", None, s_cnt + 1
            elif curr in hol_list: status, cin, h_cnt = "Holiday", None, h_cnt + 1
            else: status, cin, a_cnt = "Absent", None, a_cnt + 1
            history[m_key].append({"date": curr.strftime("%Y-%m-%d"), "check_in_time": cin, "status": status})
            curr += timedelta(days=1)

        w_days = (end_date - start_date).days + 1 - s_cnt - h_cnt
        percentage = f"{round(((p_cnt+l_cnt)/w_days)*100,1)}%" if w_days > 0 else "0%"
        return Response({
            "period": period_label,
            "teacher_profile": teacher_profile_data,
            "summary": {
                "total_days_passed": (end_date-start_date).days+1, "sundays": s_cnt, "holidays": h_cnt, 
                "holiday_dates": [h.strftime("%Y-%m-%d") for h in hol_list],
                "actual_working_days": max(0, w_days), "present": p_cnt, "late": l_cnt, "absent": a_cnt,
                "percentage": percentage,
                "attendance_percentage": percentage,
                "joining_date": joining_date.isoformat() if joining_date else None,
                "calculated_from": start_date.isoformat(),
                "message": "Attendance is calculated from the teacher joining date." if joining_date and requested_start_date < joining_date else ""
            },
            "history": history
        }, 200)

# ==========================================
# VIEW 3: ADMIN VIEW STAFF HISTORY
# ==========================================
class AdminStaffHistoryView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        s_id = request.query_params.get('staff_id')
        staff_profile = get_object_or_404(NonTeachingStaff, staff_id=s_id)
        profile_image_url = None
        if getattr(staff_profile, 'profile_image', None):
            try:
                profile_image_url = staff_profile.profile_image.url
            except Exception:
                profile_image_url = staff_profile.profile_image.url

        staff_profile_data = {
            "staff_id": staff_profile.staff_id,
            "staff_name": staff_profile.name,
            "staff_role": staff_profile.role,
            "profile_image": profile_image_url
        }
        param_date, param_month, param_year = request.query_params.get('date'), request.query_params.get('month'), request.query_params.get('year')
        today = date.today()
        year = int(param_year) if param_year else today.year

        if param_date:
            try:
                record = StaffAttendance.objects.get(staff=staff_profile, date=param_date)
                return Response({
                    "period": "Single Date",
                    "staff_profile": staff_profile_data,
                    "data": StaffAttendanceHistorySerializer(record).data
                }, 200)
            except StaffAttendance.DoesNotExist:
                return Response({
                    "period": "Single Date",
                    "staff_profile": staff_profile_data,
                    "date": param_date,
                    "status": "Absent"
                }, 200)

        if param_year and not param_month:
            start_date, end_date = date(year, 1, 1), (today if year == today.year else date(year, 12, 31))
            period_label = f"Year {year}"
        else:
            month = int(param_month) if param_month else today.month
            _, dim = calendar.monthrange(year, month)
            start_date, end_date = date(year, month, 1), date(year, month, today.day if (year==today.year and month==today.month) else dim)
            period_label = f"{calendar.month_name[month]} {year}"

        records = StaffAttendance.objects.filter(staff=staff_profile, date__range=[start_date, end_date])
        att_map = {r.date: r for r in records}
        hols = Holiday.objects.filter(date__range=[start_date, end_date]).filter(Q(applicable_for='everyone') | Q(applicable_for='staff'))
        hol_list = list(hols.values_list('date', flat=True))

        history, s_cnt, h_cnt, p_cnt, l_cnt, a_cnt = {}, 0, 0, 0, 0, 0
        curr = start_date
        while curr <= end_date:
            m_key = str(curr.month)
            if m_key not in history: history[m_key] = []
            rec = att_map.get(curr)
            if rec:
                status, cin = rec.status, rec.check_in_time.strftime("%H:%M:%S")
                if status == 'Present': p_cnt += 1
                else: l_cnt += 1
            elif curr.weekday() == 6: status, cin, s_cnt = "Sunday", None, s_cnt + 1
            elif curr in hol_list: status, cin, h_cnt = "Holiday", None, h_cnt + 1
            else: status, cin, a_cnt = "Absent", None, a_cnt + 1
            history[m_key].append({"date": curr.strftime("%Y-%m-%d"), "check_in_time": cin, "status": status})
            curr += timedelta(days=1)

        w_days = (end_date - start_date).days + 1 - s_cnt - h_cnt
        return Response({
            "period": period_label,
            "staff_profile": staff_profile_data,
            "summary": {
                "total_days_passed": (end_date-start_date).days+1, "sundays": s_cnt, "holidays": h_cnt, 
                "holiday_dates": [h.strftime("%Y-%m-%d") for h in hol_list],
                "actual_working_days": max(0, w_days), "present": p_cnt, "late": l_cnt, "absent": a_cnt,
                "percentage": f"{round(((p_cnt+l_cnt)/w_days)*100,1)}%" if w_days > 0 else "0%"
            },
            "history": history
        }, 200)

# ==========================================
# VIEW 4: ADMIN VIEW TEACHER HISTORY
# ==========================================
class AdminTeacherHistoryView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        t_id = request.query_params.get('teacher_id')
        teacher_profile = get_object_or_404(Teacher, teacher_id=t_id)
        profile_image_url = None
        if getattr(teacher_profile, 'profile_image', None):
            try:
                profile_image_url = teacher_profile.profile_image.url
            except Exception:
                profile_image_url = teacher_profile.profile_image.url

        teacher_profile_data = {
            "teacher_id": teacher_profile.teacher_id,
            "teacher_name": teacher_profile.name,
            "department": teacher_profile.department,
            "profile_image": profile_image_url
        }
        param_date, param_month, param_year = request.query_params.get('date'), request.query_params.get('month'), request.query_params.get('year')
        today = date.today()
        year = int(param_year) if param_year else today.year

        if param_date:
            try:
                record = TeacherAttendance.objects.get(teacher=teacher_profile, date=param_date)
                return Response({
                    "period": "Single Date",
                    "teacher_profile": teacher_profile_data,
                    "data": TeacherAttendanceHistorySerializer(record).data
                }, 200)
            except TeacherAttendance.DoesNotExist:
                return Response({
                    "period": "Single Date",
                    "teacher_profile": teacher_profile_data,
                    "date": param_date,
                    "status": "Absent"
                }, 200)

        if param_year and not param_month:
            start_date, end_date = date(year, 1, 1), (today if year == today.year else date(year, 12, 31))
            period_label = f"Year {year}"
        else:
            month = int(param_month) if param_month else today.month
            _, dim = calendar.monthrange(year, month)
            start_date, end_date = date(year, month, 1), date(year, month, today.day if (year==today.year and month==today.month) else dim)
            period_label = f"{calendar.month_name[month]} {year}"

        records = TeacherAttendance.objects.filter(teacher=teacher_profile, date__range=[start_date, end_date])
        att_map = {r.date: r for r in records}
        hols = Holiday.objects.filter(date__range=[start_date, end_date]).filter(Q(applicable_for='everyone') | Q(applicable_for='teachers'))
        hol_list = list(hols.values_list('date', flat=True))

        history, s_cnt, h_cnt, p_cnt, l_cnt, a_cnt = {}, 0, 0, 0, 0, 0
        curr = start_date
        while curr <= end_date:
            m_key = str(curr.month)
            if m_key not in history: history[m_key] = []
            rec = att_map.get(curr)
            if rec:
                status, cin = rec.status, rec.check_in_time.strftime("%H:%M:%S")
                if status == 'Present': p_cnt += 1
                else: l_cnt += 1
            elif curr.weekday() == 6: status, cin, s_cnt = "Sunday", None, s_cnt + 1
            elif curr in hol_list: status, cin, h_cnt = "Holiday", None, h_cnt + 1
            else: status, cin, a_cnt = "Absent", None, a_cnt + 1
            history[m_key].append({"date": curr.strftime("%Y-%m-%d"), "check_in_time": cin, "status": status})
            curr += timedelta(days=1)

        w_days = (end_date - start_date).days + 1 - s_cnt - h_cnt
        return Response({
            "period": period_label,
            "teacher_profile": teacher_profile_data,
            "summary": {
                "total_days_passed": (end_date-start_date).days+1, "sundays": s_cnt, "holidays": h_cnt, 
                "holiday_dates": [h.strftime("%Y-%m-%d") for h in hol_list],
                "actual_working_days": max(0, w_days), "present": p_cnt, "late": l_cnt, "absent": a_cnt,
                "percentage": f"{round(((p_cnt+l_cnt)/w_days)*100,1)}%" if w_days > 0 else "0%"
            },
            "history": history
        }, 200)

# ==========================================
# VIEW 5: TEACHER MARK ATTENDANCE
# ==========================================
class MarkTeacherAttendanceView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        if not hasattr(request.user, 'teacher_profile'): return Response({"error": "Unauthorized"}, 403)
        
        today = timezone.now().date()

        # Holiday Block Logic
        if Holiday.objects.filter(date=today).filter(
            Q(applicable_for='everyone') | Q(applicable_for='teachers')
        ).exists():
            return Response({"error": "Today is a holiday. Attendance not allowed."}, 400)

        is_allowed, result = check_geofence_and_time(request.data.get('latitude'), request.data.get('longitude'))
        if not is_allowed: return Response({"error": result}, 403)

        teacher = request.user.teacher_profile
        if TeacherAttendance.objects.filter(teacher=teacher, date=today).exists():
            return Response({"message": "Already Marked"}, 200)

        TeacherAttendance.objects.create(teacher=teacher, date=today, status=result)
        return Response({"message": f"Marked {result}", "status": result}, 201)


# ==========================================
# VIEW 6: STAFF MARK ATTENDANCE
# ==========================================
class MarkStaffAttendanceView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        if not hasattr(request.user, 'staff_profile'): return Response({"error": "Unauthorized"}, 403)
        
        today = timezone.now().date()

        # Holiday Block Logic
        if Holiday.objects.filter(date=today).filter(
            Q(applicable_for='everyone') | Q(applicable_for='staff')
        ).exists():
            return Response({"error": "Today is a holiday. Attendance not allowed."}, 400)

        is_allowed, result = check_geofence_and_time(request.data.get('latitude'), request.data.get('longitude'))
        if not is_allowed: return Response({"error": result}, 403)

        staff = request.user.staff_profile
        if StaffAttendance.objects.filter(staff=staff, date=today).exists():
            return Response({"message": "Already Marked"}, 200)

        StaffAttendance.objects.create(staff=staff, date=today, status=result)
        return Response({"message": f"Marked {result}", "status": result}, 201)


# ==========================================
# VIEW 7: ADMIN DAILY REPORT - TEACHER
# ==========================================
class AdminDailyTeacherReportView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        target_date = request.query_params.get('date')
        status_filter = request.query_params.get('status') 

        if not target_date: return Response({"error": "Date required"}, 400)

        all_teachers = Teacher.objects.all()
        att_map = {r.teacher.id: r for r in TeacherAttendance.objects.filter(date=target_date)}
        data = []

        for t in all_teachers:
            record = att_map.get(t.id)
            status = record.status if record else "Absent"
            if status_filter and status != status_filter: continue
            
            data.append({
                "id": t.teacher_id, 
                "name": t.name, 
                "status": status, 
                "check_in_time": record.check_in_time if record else None
            })
        
        return Response({"count": len(data), "data": DailyTeacherStatusSerializer(data, many=True).data})


class AdminDailyTeacherReportPagination(PageNumberPagination):
    page_size = 10
    page_size_query_param = "page_size"
    max_page_size = 100


class AdminDailyTeacherReportPaginatedView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        target_date = request.query_params.get('date')
        status_filter = (request.query_params.get('status') or '').strip()
        search = (request.query_params.get('search') or '').strip().lower()

        if not target_date:
            return Response({"error": "Date required"}, 400)

        all_teachers = Teacher.objects.all()
        att_map = {r.teacher.id: r for r in TeacherAttendance.objects.filter(date=target_date)}
        data = []

        for t in all_teachers:
            record = att_map.get(t.id)
            teacher_status = record.status if record else "Absent"

            if status_filter and status_filter.lower() != 'all' and teacher_status.lower() != status_filter.lower():
                continue

            if search and (search not in (t.name or '').lower() and search not in (t.teacher_id or '').lower()):
                continue

            data.append({
                "id": t.teacher_id,
                "name": t.name,
                "status": teacher_status,
                "check_in_time": record.check_in_time if record else None
            })

        total_filtered = len(data)
        present_count = sum(1 for item in data if item["status"] == "Present")
        late_count = sum(1 for item in data if item["status"] == "Late")
        absent_count = sum(1 for item in data if item["status"] == "Absent")
        present_percentage = round((present_count / total_filtered) * 100, 1) if total_filtered > 0 else 0.0

        paginator = AdminDailyTeacherReportPagination()
        page_data = paginator.paginate_queryset(data, request, view=self)
        serialized = DailyTeacherStatusSerializer(page_data, many=True).data

        return Response({
            "count": total_filtered,
            "next": paginator.get_next_link(),
            "previous": paginator.get_previous_link(),
            "results": serialized,
            "summary": {
                "total": total_filtered,
                "present": present_count,
                "late": late_count,
                "absent": absent_count,
                "present_percentage": present_percentage,
            },
        })


# ==========================================
# VIEW 8: ADMIN DAILY REPORT - STAFF
# ==========================================
class AdminDailyStaffReportView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        target_date = request.query_params.get('date')
        status_filter = request.query_params.get('status')
        role_filter = request.query_params.get('role')

        if not target_date: return Response({"error": "Date required"}, 400)

        all_staff = NonTeachingStaff.objects.all()
        if role_filter: all_staff = all_staff.filter(role__iexact=role_filter)
        att_map = {r.staff.id: r for r in StaffAttendance.objects.filter(date=target_date)}
        data = []

        for s in all_staff:
            record = att_map.get(s.id)
            status = record.status if record else "Absent"
            if status_filter and status != status_filter: continue

            data.append({
                "id": s.staff_id, 
                "name": s.name, 
                "role": s.role, 
                "status": status, 
                "check_in_time": record.check_in_time if record else None
            })
        return Response({"count": len(data), "data": DailyStaffStatusSerializer(data, many=True).data})


class AdminDailyStaffReportPagination(PageNumberPagination):
    page_size = 10
    page_size_query_param = "page_size"
    max_page_size = 100


class AdminDailyStaffReportPaginatedView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        target_date = request.query_params.get('date')
        status_filter = (request.query_params.get('status') or '').strip()
        role_filter = (request.query_params.get('role') or '').strip()
        search = (request.query_params.get('search') or '').strip().lower()

        if not target_date:
            return Response({"error": "Date required"}, 400)

        all_staff = NonTeachingStaff.objects.all()
        if role_filter and role_filter.lower() != 'all':
            all_staff = all_staff.filter(role__iexact=role_filter)

        att_map = {r.staff.id: r for r in StaffAttendance.objects.filter(date=target_date)}
        data = []

        for s in all_staff:
            record = att_map.get(s.id)
            staff_status = record.status if record else "Absent"

            if status_filter and status_filter.lower() != 'all' and staff_status.lower() != status_filter.lower():
                continue

            if search and (
                search not in (s.name or '').lower()
                and search not in (s.staff_id or '').lower()
                and search not in (s.role or '').lower()
            ):
                continue

            data.append({
                "id": s.staff_id,
                "name": s.name,
                "role": s.role,
                "status": staff_status,
                "check_in_time": record.check_in_time if record else None
            })

        total_filtered = len(data)
        present_count = sum(1 for item in data if item["status"] == "Present")
        late_count = sum(1 for item in data if item["status"] == "Late")
        absent_count = sum(1 for item in data if item["status"] == "Absent")
        present_percentage = round((present_count / total_filtered) * 100, 1) if total_filtered > 0 else 0.0

        role_summary = {}
        for item in data:
            role = item.get("role") or "unknown"
            if role not in role_summary:
                role_summary[role] = {
                    "total": 0,
                    "present": 0,
                    "late": 0,
                    "absent": 0,
                }
            role_summary[role]["total"] += 1
            status = item.get("status")
            if status == "Present":
                role_summary[role]["present"] += 1
            elif status == "Late":
                role_summary[role]["late"] += 1
            elif status == "Absent":
                role_summary[role]["absent"] += 1

        paginator = AdminDailyStaffReportPagination()
        page_data = paginator.paginate_queryset(data, request, view=self)
        serialized = DailyStaffStatusSerializer(page_data, many=True).data

        return Response({
            "count": total_filtered,
            "next": paginator.get_next_link(),
            "previous": paginator.get_previous_link(),
            "results": serialized,
            "summary": {
                "total": total_filtered,
                "present": present_count,
                "late": late_count,
                "absent": absent_count,
                "present_percentage": present_percentage,
            },
            "role_summary": role_summary,
        })


# ==========================================
# VIEW 9: ADMIN CONFIG CRUD
# ==========================================
class AttendanceConfigView(APIView):
    permission_classes = [IsAuthenticated, IsOwnerAdminOrAdminStaff]

    def get(self, request):
        config = AttendanceConfig.objects.first()
        if not config: return Response({"message": "Not configured"}, 200)
        return Response(AttendanceConfigSerializer(config).data)

    def post(self, request):
        config = AttendanceConfig.objects.first()
        serializer = AttendanceConfigSerializer(config, data=request.data) if config else AttendanceConfigSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, 200)
        return Response(serializer.errors, 400)

    def put(self, request):
        return self.post(request)

    def delete(self, request):
        AttendanceConfig.objects.all().delete()
        return Response({"message": "Config Deleted"}, 200)


### SIVA BRO ATTENDANCE VIEWS# # # # ### #


class ClassWiseAttendanceReportView(APIView):
    """
    GET: Get attendance summary for all classes on a specific date.
    Usage: /api/attendance/class-report/?date=2025-01-01
    """
    permission_classes = [IsAuthenticated, IsAdmin]
    
    def get(self, request):
        date_str = request.query_params.get('date')
        if not date_str:
            # Default to today
            date_str = datetime.now().date().isoformat()
        
        try:
            date_obj = datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            return Response({"error": "Invalid date format. Use YYYY-MM-DD"}, status=400)
        
        # Get all classes (standards)
        classes = Standard.objects.all().order_by('name')
        class_report = []
        
        for standard in classes:
            # Get all sections for this class
            sections = Section.objects.filter(standard=standard)
            
            # Get total students in this class across all sections
            total_students = Enrollment.objects.filter(
                section__in=sections,
                is_active=True
            ).count()
            
            # Get attendance counts for this date across all sections of this class
            attendance_counts = Attendance.objects.filter(
                enrollment__section__standard=standard,
                date=date_obj
            ).aggregate(
                present=Count('id', filter=Q(status__iexact='present')),
                absent=Count('id', filter=Q(status__iexact='absent')),
                late=Count('id', filter=Q(status__iexact='late'))
            )
            
            present = attendance_counts['present'] or 0
            absent = attendance_counts['absent'] or 0
            late = attendance_counts['late'] or 0
            
            # Calculate attendance percentage
            total_marked = present + absent + late
            attendance_percentage = (present / total_marked * 100) if total_marked > 0 else 0
            
            class_report.append({
                "class_id": standard.id,
                "class_name": standard.name,
                "total_students": total_students,
                "present": present,
                "absent": absent,
                "late": late,
                "attendance_percentage": round(attendance_percentage, 1),
                "section_count": sections.count()
            })
        
        # Sort classes numerically (handles 'LKG', 'UKG', '1', '2', etc.)
        def sort_key(item):
            name = item['class_name']
            if name.isdigit():
                return int(name)
            # Handle special class names
            if name.upper() == 'LKG':
                return -2
            if name.upper() == 'UKG':
                return -1
            if name.upper() == 'PRE-KG':
                return -3
            return 999  # Put other non-numeric classes at the end
        
        sorted_report = sorted(class_report, key=sort_key)
        
        # Calculate school-wide totals
        total_marked_all = sum(item['present'] + item['absent'] + item['late'] for item in sorted_report)
        total_students_all = sum(item['total_students'] for item in sorted_report)
        school_percentage = (sum(item['present'] for item in sorted_report) / total_marked_all * 100) if total_marked_all > 0 else 0
        
        school_totals = {
            "total_classes": len(sorted_report),
            "total_students": total_students_all,
            "total_present": sum(item['present'] for item in sorted_report),
            "total_absent": sum(item['absent'] for item in sorted_report),
            "total_late": sum(item['late'] for item in sorted_report),
            "total_marked": total_marked_all,
            "attendance_percentage": round(school_percentage, 1)
        }
        
        return Response({
            "status": 200,
            "date": date_str,
            "school_summary": school_totals,
            "class_report": sorted_report
        })
    


class TodayClassAttendanceReportView(APIView):
    """
    GET: Get today's attendance summary for all classes.
    Request: GET /api/attendance/today-class-report/
    Response: Exact format as specified - class-wise attendance for today only
    """
    permission_classes = [IsAuthenticated]  # Changed from AllowAny for security
    
    def get(self, request):
        # Get today's date
        print("it comming")
        today_date = datetime.now().date()
        date_str = today_date.isoformat()
        
        try:
            # Get all classes (standards)
            classes = Standard.objects.all().order_by('name')
            
            if not classes.exists():
                return Response({
                    "status": 404,
                    "message": "No classes found",
                    "date": date_str,
                    "school_summary": {
                        "total_classes": 0,
                        "total_students": 0,
                        "total_present": 0,
                        "total_absent": 0,
                        "total_late": 0,
                        "total_marked": 0,
                        "attendance_percentage": 0.0
                    },
                    "class_report": []
                }, status=200)
            
            class_report = []
            total_present_all = 0
            total_absent_all = 0
            total_late_all = 0
            total_marked_all = 0
            total_students_all = 0
            
            for standard in classes:
                try:
                    # Get all sections for this class
                    sections = Section.objects.filter(standard=standard)
                    
                    if not sections.exists():
                        continue
                    
                    # Get total active students in this class across all sections
                    total_students = Enrollment.objects.filter(
                        section__in=sections,
                        is_active=True
                    ).count()
                    
                    # Get attendance counts for today
                    attendance_counts = Attendance.objects.filter(
                        enrollment__section__standard=standard,
                        date=today_date
                    ).aggregate(
                        present=Count('id', filter=Q(status__iexact='present')),
                        absent=Count('id', filter=Q(status__iexact='absent')),
                        late=Count('id', filter=Q(status__iexact='late'))
                    )
                    
                    present = attendance_counts['present'] or 0
                    absent = attendance_counts['absent'] or 0
                    late = attendance_counts['late'] or 0
                    total_marked = present + absent + late
                    
                    # Calculate attendance percentage
                    if total_students > 0:
                        attendance_percentage = round((present / total_students * 100), 1) if total_students > 0 else 0.0
                    else:
                        attendance_percentage = 0.0
                    
                    # Add to totals
                    total_present_all += present
                    total_absent_all += absent
                    total_late_all += late
                    total_marked_all += total_marked
                    total_students_all += total_students
                    
                    class_report.append({
                        "class_id": standard.id,
                        "class_name": standard.name,
                        "total_students": total_students,
                        "present": present,
                        "absent": absent,
                        "late": late,
                        "attendance_percentage": attendance_percentage,
                        "section_count": sections.count()
                    })
                    
                except Exception as e:
                    # Log error but continue processing other classes
                    print(f"Error processing class {standard.name}: {str(e)}")
                    continue
            
            # Sort classes: Pre-KG, LKG, UKG, then 1,2,3...
            def sort_key(item):
                name = item['class_name']
                if name.upper() == 'PRE-KG':
                    return -3
                if name.upper() == 'LKG':
                    return -2
                if name.upper() == 'UKG':
                    return -1
                if name.isdigit():
                    return int(name)
                return 999
            
            sorted_report = sorted(class_report, key=sort_key)
            
            # Calculate school-wide totals
            school_percentage = round(
                (total_present_all / total_students_all * 100) if total_students_all > 0 else 0.0, 
                1
            )
            
            school_totals = {
                "total_classes": len(sorted_report),
                "total_students": total_students_all,
                "total_present": total_present_all,
                "total_absent": total_absent_all,
                "total_late": total_late_all,
                "total_marked": total_marked_all,
                "attendance_percentage": school_percentage
            }

            print(school_totals)
            
            return Response({
                "status": 200,
                "date": date_str,
                "school_summary": school_totals,
                "class_report": sorted_report
            }, status=200)
            
        except Standard.DoesNotExist:
            return Response({
                "status": 404,
                "error": "Standards configuration not found",
                "message": "No class standards configured in the system"
            }, status=404)
            
        except Exception as e:
            return Response({
                "status": 500,
                "error": "Internal server error",
                "message": str(e)
            }, status=500)



class SectionWiseAttendanceDetailView(APIView):
    """
    GET: Get detailed attendance for a specific class (all sections)
    Usage: /api/attendance/class-detail/?class=10&date=2025-01-01
    """
    permission_classes = [IsAuthenticated, IsAdmin]
    
    def get(self, request):
        class_name = (request.query_params.get('class') or "").strip()
        date_str = request.query_params.get('date')
        
        if not class_name:
            return Response({"error": "Class parameter is required"}, status=400)
        
        if not date_str:
            # Default to today
            date_str = datetime.now().date().isoformat()
        
        try:
            date_obj = datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            return Response({"error": "Invalid date format. Use YYYY-MM-DD"}, status=400)
        
        try:
            standard = Standard.objects.get(name=class_name)
        except Standard.DoesNotExist:
            return Response({"error": f"Class '{class_name}' not found"}, status=404)
        
        sections = Section.objects.filter(standard=standard).order_by('name')
        current_academic_year = AcademicYear.objects.filter(is_current=True).first()
        current_academic_year_data = {
            "id": current_academic_year.id,
            "name": current_academic_year.name
        } if current_academic_year else {
            "id": "",
            "name": "",
            "message": "No active academic year configured"
        }

        if not sections.exists():
            return Response({
                "status": 200,
                "class": class_name,
                "date": date_str,
                "current_academic_year": current_academic_year_data,
                "class_summary": {
                    "total_sections": 0,
                    "total_students": 0,
                    "total_present": 0,
                    "total_absent": 0,
                    "total_late": 0,
                    "total_unmarked": 0,
                    "total_marked": 0,
                    "overall_attendance": 0.0
                },
                "section_details": []
            })

        section_details = []
        today_date = timezone.localdate()

        def get_day_info(target_date):
            is_sunday = target_date.weekday() == 6
            is_holiday = Holiday.objects.filter(date=target_date).filter(
                Q(applicable_for='everyone') | Q(applicable_for='students_only')
            ).exists()
            return is_sunday, is_holiday

        def normalize_status(raw_status, is_sunday, is_holiday):
            if is_sunday:
                return "Sunday"
            if is_holiday:
                return "Holiday"
            if not raw_status:
                return "Not Marked"

            status_key = str(raw_status).strip().lower()
            if status_key == "present":
                return "Present"
            if status_key == "absent":
                return "Absent"
            if status_key == "late":
                return "Late"
            return "Not Marked"
        
        for section in sections:
            # Get total students in this section
            enrollments = Enrollment.objects.filter(
                section=section,
                is_active=True
            ).select_related('student').order_by('student__student_name')
            total_students = enrollments.count()
            enrollment_ids = list(enrollments.values_list('id', flat=True))
            
            # Get attendance for this section on the given date
            # FIXED: Filter through enrollment__section instead of section
            attendance_counts = Attendance.objects.filter(
    enrollment__section=section,  # <-- This is correct
    date=date_obj
).aggregate(
    present=Count('id', filter=Q(status__iexact='present')),
    absent=Count('id', filter=Q(status__iexact='absent')),
    late=Count('id', filter=Q(status__iexact='late'))
)
            
            present = attendance_counts['present'] or 0
            absent = attendance_counts['absent'] or 0
            late = attendance_counts['late'] or 0
            total_marked = present + absent + late
            attendance_percentage = (present / total_marked * 100) if total_marked > 0 else 0
            unmarked = total_students - total_marked
            selected_is_sunday, selected_is_holiday = get_day_info(date_obj)
            if selected_is_sunday or selected_is_holiday:
                present = 0
                absent = 0
                late = 0
                total_marked = 0
                unmarked = 0
                attendance_percentage = 0

            attendance_for_selected_date = Attendance.objects.filter(
                enrollment_id__in=enrollment_ids,
                date=date_obj
            ).values('enrollment_id', 'status')
            selected_date_status_map = {
                row['enrollment_id']: row['status']
                for row in attendance_for_selected_date
            }

            if today_date == date_obj:
                today_status_map = selected_date_status_map
            else:
                attendance_for_today = Attendance.objects.filter(
                    enrollment_id__in=enrollment_ids,
                    date=today_date
                ).values('enrollment_id', 'status')
                today_status_map = {
                    row['enrollment_id']: row['status']
                    for row in attendance_for_today
                }
            today_is_sunday, today_is_holiday = get_day_info(today_date)

            students = []
            today_present = 0
            today_absent = 0
            today_late = 0
            for enrollment in enrollments:
                selected_status = normalize_status(
                    selected_date_status_map.get(enrollment.id),
                    selected_is_sunday,
                    selected_is_holiday
                )
                today_status = normalize_status(
                    today_status_map.get(enrollment.id),
                    today_is_sunday,
                    today_is_holiday
                )

                if today_status:
                    normalized_today = today_status.lower()
                    if normalized_today == 'present':
                        today_present += 1
                    elif normalized_today == 'absent':
                        today_absent += 1
                    elif normalized_today == 'late':
                        today_late += 1

                students.append({
                    "student_id": enrollment.student.student_id,
                    "student_name": enrollment.student.student_name,
                    "roll_no": getattr(enrollment.student, 'roll_no', "") or "",
                    "attendance_status": selected_status,
                    "today_attendance": today_status
                })

            today_total_marked = today_present + today_absent + today_late
            today_unmarked = total_students - today_total_marked
            today_percentage = (today_present / today_total_marked * 100) if today_total_marked > 0 else 0
            today_day_status = "Sunday" if today_is_sunday else "Holiday" if today_is_holiday else "Working Day"
            if today_is_sunday or today_is_holiday:
                today_unmarked = 0
                today_percentage = 0
            
            # Get class teacher from ClassTeacher model
            class_teacher_record = {
                "name": "Not Assigned",
                "teacher_id": "",
                "message": "Class teacher not assigned"
            }
            try:
                active_year = AcademicYear.objects.get(is_current=True)
                ct = ClassTeacher.objects.get(
                    section=section,
                    academic_year=active_year
                )
                class_teacher_record = {
                    "name": (ct.teacher.name if ct.teacher else "") or "Not Assigned",
                    "teacher_id": (ct.teacher.teacher_id if ct.teacher else "") or "",
                    "message": "Assigned"
                }
            except AcademicYear.DoesNotExist:
                class_teacher_record = {
                    "name": "Not Assigned",
                    "teacher_id": "",
                    "message": "Active academic year not configured"
                }
            except ClassTeacher.DoesNotExist:
                class_teacher_record = {
                    "name": "Not Assigned",
                    "teacher_id": "",
                    "message": "Class teacher not assigned"
                }
            
            section_details.append({
                "section_id": section.id,
                "section_name": section.name,
                "total_students": total_students,
                "present": present,
                "absent": absent,
                "late": late,
                "unmarked": unmarked if unmarked > 0 else 0,
                "attendance_percentage": round(attendance_percentage, 1),
                "class_teacher": class_teacher_record,
                "today_attendance": {
                    "date": today_date.isoformat(),
                    "day_status": today_day_status,
                    "present": today_present,
                    "absent": today_absent,
                    "late": today_late,
                    "unmarked": today_unmarked if today_unmarked > 0 else 0,
                    "attendance_percentage": round(today_percentage, 1)
                },
                "students": students
            })
        
        # Calculate class totals
        class_totals = {
            "total_sections": sections.count(),
            "total_students": sum(item['total_students'] for item in section_details),
            "total_present": sum(item['present'] for item in section_details),
            "total_absent": sum(item['absent'] for item in section_details),
            "total_late": sum(item['late'] for item in section_details),
            "total_unmarked": sum(item['unmarked'] for item in section_details),
            "total_marked": sum(item['present'] + item['absent'] + item['late'] for item in section_details),
            "overall_attendance": round(
                (sum(item['present'] for item in section_details) / 
                 max(1, sum(item['present'] + item['absent'] + item['late'] for item in section_details))) * 100, 
                1
            )
        }
        
        return Response({
            "status": 200,
            "class": class_name,
            "date": date_str,
            "current_academic_year": current_academic_year_data,
            "class_summary": class_totals,
            "section_details": section_details
        })
    

class AdminAttendanceOverviewView(APIView):
    """
    GET: Get attendance overview for chart display with date range and class filters.
    Query Params:
        - period: 'this_week', 'past_week', 'past_two_weeks', 'this_month', 'past_month'
        - class_name: (optional) Filter by specific class (e.g., '10'), or 'all' for all classes
    Returns:
        - Daily counts of present/absent/late for the period
        - Class-wise breakdown if class_name='all'
        - Excludes Sundays and holidays
    """
    permission_classes = [IsAuthenticated, IsAdmin]
    
    def get_date_range(self, period):
        """Calculate date range based on period parameter"""
        today = timezone.now().date()
        
        if period == 'this_week':
            # Monday to Sunday of current week
            start_date = today - timedelta(days=today.weekday())  # Monday
            end_date = start_date + timedelta(days=6)  # Sunday
            return start_date, end_date
            
        elif period == 'past_week':
            # Previous Monday to Sunday
            start_date = today - timedelta(days=today.weekday() + 7)  # Previous Monday
            end_date = start_date + timedelta(days=6)  # Previous Sunday
            return start_date, end_date
            
        elif period == 'past_two_weeks':
            # Last 14 days (excluding today)
            end_date = today - timedelta(days=1)
            start_date = end_date - timedelta(days=13)
            return start_date, end_date
            
        elif period == 'this_month':
            # 1st to last day of current month
            start_date = today.replace(day=1)
            # Get last day of month
            if today.month == 12:
                end_date = today.replace(year=today.year + 1, month=1, day=1) - timedelta(days=1)
            else:
                end_date = today.replace(month=today.month + 1, day=1) - timedelta(days=1)
            return start_date, end_date
            
        elif period == 'past_month':
            # Complete previous month
            if today.month == 1:
                prev_month = 12
                prev_year = today.year - 1
            else:
                prev_month = today.month - 1
                prev_year = today.year
            
            start_date = today.replace(year=prev_year, month=prev_month, day=1)
            # Get last day of previous month
            if prev_month == 12:
                end_date = today.replace(year=prev_year + 1, month=1, day=1) - timedelta(days=1)
            else:
                end_date = today.replace(year=prev_year, month=prev_month + 1, day=1) - timedelta(days=1)
            return start_date, end_date
        
        # Default: last 7 days
        end_date = today - timedelta(days=1)
        start_date = end_date - timedelta(days=6)
        return start_date, end_date
    
    def is_working_day(self, date_obj):
        """Check if a date is a working day (not Sunday and not holiday)"""
        # Check if Sunday
        if date_obj.weekday() == 6:  # Sunday is 6
            return False
        
        # Check if holiday
        if Holiday.objects.filter(date=date_obj).filter(
            Q(applicable_for='everyone') | Q(applicable_for='students')
        ).exists():
            return False
            
        return True
    
    def get(self, request):
        # Get query parameters
        period = request.query_params.get('period', 'this_week')
        class_name = request.query_params.get('class_name', 'all')
        
        # Validate period parameter
        valid_periods = ['this_week', 'past_week', 'past_two_weeks', 'this_month', 'past_month']
        if period not in valid_periods:
            return Response({"error": f"Invalid period. Choose from: {', '.join(valid_periods)}"}, status=400)
        
        # Get date range
        start_date, end_date = self.get_date_range(period)
        
        # Get all working days in the range (exclude Sundays and holidays)
        working_days = []
        current_date = start_date
        while current_date <= end_date:
            if self.is_working_day(current_date):
                working_days.append(current_date)
            current_date += timedelta(days=1)
        
        if not working_days:
            return Response({
                "message": "No working days in the selected period",
                "period": period,
                "start_date": start_date,
                "end_date": end_date
            }, status=200)
        
        # Initialize response structure
        response_data = {
            "period": period,
            "date_range": {
                "start": start_date.isoformat(),
                "end": end_date.isoformat()
            },
            "total_working_days": len(working_days),
            "data": {}
        }
        
        # CASE 1: Specific class (not 'all')
        if class_name != 'all':
            try:
                standard = Standard.objects.get(name=class_name)
            except Standard.DoesNotExist:
                return Response({"error": f"Class '{class_name}' not found"}, status=404)
            
            # Get all sections for this class
            sections = Section.objects.filter(standard=standard)
            
            # Initialize daily data structure
            daily_data = []
            
            for date_obj in working_days:
                # Get total students in this class across all sections
                total_students = Enrollment.objects.filter(
                    section__in=sections,
                    is_active=True
                ).count()
                
                # Get attendance counts for this date and class
                attendance_counts = Attendance.objects.filter(
                    enrollment__section__standard=standard,
                    date=date_obj
                ).aggregate(
                    present=Count('id', filter=Q(status__iexact='present')),
                    absent=Count('id', filter=Q(status__iexact='absent')),
                    late=Count('id', filter=Q(status__iexact='late'))
                )
                
                present = attendance_counts['present'] or 0
                absent = attendance_counts['absent'] or 0
                late = attendance_counts['late'] or 0
                total_marked = present + absent + late
                
                # Calculate attendance percentage
                attendance_percentage = (present / total_marked * 100) if total_marked > 0 else 0
                
                daily_data.append({
                    "date": date_obj.isoformat(),
                    "day": date_obj.strftime("%a"),  # Mon, Tue, etc.
                    "total_students": total_students,
                    "present": present,
                    "absent": absent,
                    "late": late,
                    "total_marked": total_marked,
                    "unmarked": total_students - total_marked,
                    "attendance_percentage": round(attendance_percentage, 1)
                })
            
            # Calculate period totals
            period_totals = {
                "total_present": sum(day['present'] for day in daily_data),
                "total_absent": sum(day['absent'] for day in daily_data),
                "total_late": sum(day['late'] for day in daily_data),
                "total_marked": sum(day['total_marked'] for day in daily_data),
                "average_attendance": round(
                    sum(day['attendance_percentage'] for day in daily_data) / len(daily_data) 
                    if daily_data else 0, 1
                )
            }
            
            response_data["data"] = {
                "class": class_name,
                "type": "single_class",
                "daily_data": daily_data,
                "totals": period_totals
            }
            
        # CASE 2: All classes
        else:
            # Get all classes
            classes = Standard.objects.all().order_by('name')
            
            # Initialize class-wise data structure
            class_data = []
            all_daily_totals = defaultdict(lambda: {
                "present": 0,
                "absent": 0,
                "late": 0,
                "total_students": 0
            })
            
            # Helper function to sort class names numerically
            def sort_class_key(cls_name):
                if cls_name.isdigit():
                    return int(cls_name)
                if cls_name.upper() == 'LKG':
                    return -2
                if cls_name.upper() == 'UKG':
                    return -1
                if cls_name.upper() == 'PRE-KG':
                    return -3
                return 999
            
            sorted_classes = sorted(classes, key=lambda x: sort_class_key(x.name))
            
            for standard in sorted_classes:
                # Get all sections for this class
                sections = Section.objects.filter(standard=standard)
                
                # Get total students in this class
                total_students = Enrollment.objects.filter(
                    section__in=sections,
                    is_active=True
                ).count()
                
                # Get attendance counts for this class across all dates
                attendance_totals = Attendance.objects.filter(
                    enrollment__section__standard=standard,
                    date__in=working_days
                ).aggregate(
                    present=Count('id', filter=Q(status__iexact='present')),
                    absent=Count('id', filter=Q(status__iexact='absent')),
                    late=Count('id', filter=Q(status__iexact='late'))
                )
                
                present = attendance_totals['present'] or 0
                absent = attendance_totals['absent'] or 0
                late = attendance_totals['late'] or 0
                total_marked = present + absent + late
                
                # Calculate attendance percentage
                attendance_percentage = (present / total_marked * 100) if total_marked > 0 else 0
                
                # Add to class data
                class_data.append({
                    "class_name": standard.name,
                    "total_students": total_students,
                    "present": present,
                    "absent": absent,
                    "late": late,
                    "total_marked": total_marked,
                    "attendance_percentage": round(attendance_percentage, 1),
                    "section_count": sections.count()
                })
                
                # Get daily breakdown for this class
                daily_breakdown = []
                for date_obj in working_days:
                    day_attendance = Attendance.objects.filter(
                        enrollment__section__standard=standard,
                        date=date_obj
                    ).aggregate(
                        present=Count('id', filter=Q(status__iexact='present')),
                        absent=Count('id', filter=Q(status__iexact='absent')),
                        late=Count('id', filter=Q(status__iexact='late'))
                    )
                    
                    day_present = day_attendance['present'] or 0
                    day_absent = day_attendance['absent'] or 0
                    day_late = day_attendance['late'] or 0
                    
                    # Add to overall daily totals
                    date_key = date_obj.isoformat()
                    all_daily_totals[date_key]["present"] += day_present
                    all_daily_totals[date_key]["absent"] += day_absent
                    all_daily_totals[date_key]["late"] += day_late
                    all_daily_totals[date_key]["total_students"] += total_students
                    
                    daily_breakdown.append({
                        "date": date_key,
                        "present": day_present,
                        "absent": day_absent,
                        "late": day_late
                    })
                
                # Add daily breakdown to class data
                class_data[-1]["daily_breakdown"] = daily_breakdown
            
            # Prepare overall daily summary
            daily_summary = []
            for date_obj in working_days:
                date_key = date_obj.isoformat()
                day_data = all_daily_totals[date_key]
                total_marked = day_data["present"] + day_data["absent"] + day_data["late"]
                
                daily_summary.append({
                    "date": date_key,
                    "day": date_obj.strftime("%a"),
                    "total_students": day_data["total_students"],
                    "present": day_data["present"],
                    "absent": day_data["absent"],
                    "late": day_data["late"],
                    "total_marked": total_marked,
                    "unmarked": day_data["total_students"] - total_marked,
                    "attendance_percentage": round(
                        (day_data["present"] / total_marked * 100) if total_marked > 0 else 0, 1
                    )
                })
            
            # Calculate overall totals
            overall_totals = {
                "total_classes": len(class_data),
                "total_students": sum(cls['total_students'] for cls in class_data),
                "total_present": sum(cls['present'] for cls in class_data),
                "total_absent": sum(cls['absent'] for cls in class_data),
                "total_late": sum(cls['late'] for cls in class_data),
                "total_marked": sum(cls['total_marked'] for cls in class_data),
                "average_attendance": round(
                    sum(cls['attendance_percentage'] for cls in class_data) / len(class_data) 
                    if class_data else 0, 1
                )
            }
            
            response_data["data"] = {
                "type": "all_classes",
                "class_summary": class_data,
                "daily_summary": daily_summary,
                "overall_totals": overall_totals
            }
        
        return Response({
            "status": 200,
            **response_data
        }, status=200)


QR_SIGNING_SALT = "attendance.qr.dynamic.v1"


def _get_client_ip(request):
    forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if forwarded_for:
        return forwarded_for.split(',')[0].strip()
    return request.META.get('REMOTE_ADDR')


def _compute_qr_slot(session, now_dt):
    # Daily static QR mode: a single slot for the full active session window.
    if now_dt < session.starts_at or now_dt > session.ends_at:
        return -1
    return 0


def _build_qr_token(session, now_dt):
    slot_no = 0
    expires_at = session.ends_at

    payload = {
        "sid": session.id,
        "slot": slot_no,
        "exp": int(expires_at.timestamp()),
        "scope": session.role_scope,
    }
    token = signing.Signer(salt=QR_SIGNING_SALT).sign_object(payload, compress=True)
    return token, slot_no, expires_at


def _load_qr_token(token):
    try:
        return signing.Signer(salt=QR_SIGNING_SALT).unsign_object(token)
    except signing.BadSignature:
        # Accept tokens created before the static QR change until they naturally expire.
        return signing.loads(token, salt=QR_SIGNING_SALT, max_age=60 * 60 * 24)


def _resolve_role_and_profile(user):
    if hasattr(user, 'teacher_profile'):
        return "teacher", user.teacher_profile
    if hasattr(user, 'staff_profile'):
        return "staff", user.staff_profile
    return None, None


def _current_mark_status():
    config = AttendanceConfig.objects.first()
    cutoff = config.late_cutoff_time if config else time(9, 0)
    return "Late" if timezone.localtime().time() > cutoff else "Present"


def _log_qr_scan_attempt(request, *, session, user, role, profile, slot_no, accepted, reason, marked_status=""):
    QRAttendanceScanLog.objects.create(
        session=session,
        user=user,
        teacher=profile if role == 'teacher' else None,
        staff=profile if role == 'staff' else None,
        token_slot=max(0, slot_no),
        accepted=accepted,
        reason=reason,
        marked_status=marked_status,
        ip_address=_get_client_ip(request),
        user_agent=request.META.get('HTTP_USER_AGENT', ''),
    )


class QRAttendanceSessionStartView(APIView):
    permission_classes = [IsAuthenticated, IsOwnerAdminOrAdminStaff]

    def post(self, request):
        serializer = QRAttendanceSessionStartSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        now_dt = timezone.now()
        local_now = timezone.localtime(now_dt)
        end_of_day = local_now.replace(hour=23, minute=59, second=59, microsecond=0)

        session = QRAttendanceSession.objects.create(
            role_scope=serializer.validated_data['role_scope'],
            starts_at=now_dt,
            ends_at=end_of_day,
            rotation_seconds=serializer.validated_data['rotation_seconds'],
            created_by=request.user,
        )

        token, slot_no, expires_at = _build_qr_token(session, now_dt)
        return Response({
            "session": QRAttendanceSessionSerializer(session).data,
            "token": token,
            "slot_no": slot_no,
            "expires_at": expires_at,
        }, status=201)


class QRAttendanceSessionTokenView(APIView):
    permission_classes = [IsAuthenticated, IsOwnerAdminOrAdminStaff]

    def get(self, request, session_id):
        session = get_object_or_404(QRAttendanceSession, id=session_id)
        now_dt = timezone.now()

        if not session.is_active or now_dt > session.ends_at:
            return Response({"error": "Session is not active"}, status=400)

        token, slot_no, expires_at = _build_qr_token(session, now_dt)
        return Response({
            "session_id": session.id,
            "token": token,
            "slot_no": slot_no,
            "expires_at": expires_at,
            "role_scope": session.role_scope,
        }, status=200)


class QRAttendanceSessionCloseView(APIView):
    permission_classes = [IsAuthenticated, IsOwnerAdminOrAdminStaff]

    def post(self, request, session_id):
        session = get_object_or_404(QRAttendanceSession, id=session_id)
        if not session.is_active:
            return Response({"message": "Session already closed"}, status=200)

        session.is_active = False
        session.closed_at = timezone.now()
        session.save(update_fields=['is_active', 'closed_at', 'updated_at'])
        return Response({"message": "Session closed"}, status=200)


class QRAttendanceScanMarkView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        serializer = QRAttendanceScanSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        role, profile = _resolve_role_and_profile(request.user)
        if not role:
            return Response({"error": "Only teachers/staff can mark QR attendance"}, status=403)

        token = serializer.validated_data['token']
        try:
            payload = _load_qr_token(token)
        except signing.BadSignature:
            return Response({"error": "Invalid QR token"}, status=400)

        session_id = payload.get("sid")
        slot_no = int(payload.get("slot", -1))
        token_exp = int(payload.get("exp", 0))

        session = get_object_or_404(QRAttendanceSession, id=session_id)
        now_dt = timezone.now()

        if not session.is_active or now_dt > session.ends_at:
            _log_qr_scan_attempt(
                request,
                session=session,
                user=request.user,
                role=role,
                profile=profile,
                slot_no=slot_no,
                accepted=False,
                reason="Session closed or expired",
            )
            return Response({"error": "Session is not active"}, status=400)

        if int(now_dt.timestamp()) > token_exp:
            _log_qr_scan_attempt(
                request,
                session=session,
                user=request.user,
                role=role,
                profile=profile,
                slot_no=slot_no,
                accepted=False,
                reason="QR token expired",
            )
            return Response({"error": "QR expired. Scan the latest QR."}, status=400)

        if session.role_scope != 'both' and session.role_scope != role:
            _log_qr_scan_attempt(
                request,
                session=session,
                user=request.user,
                role=role,
                profile=profile,
                slot_no=slot_no,
                accepted=False,
                reason="Role not allowed for session",
            )
            return Response({"error": "This QR session is not for your role"}, status=403)

        today = timezone.localdate()

        if role == 'teacher':
            is_holiday = Holiday.objects.filter(date=today).filter(
                Q(applicable_for='everyone') | Q(applicable_for='teachers')
            ).exists()
            if is_holiday:
                _log_qr_scan_attempt(
                    request,
                    session=session,
                    user=request.user,
                    role=role,
                    profile=profile,
                    slot_no=slot_no,
                    accepted=False,
                    reason="Holiday",
                )
                return Response({"error": "Today is a holiday. Attendance not allowed."}, status=400)

            if TeacherAttendance.objects.filter(teacher=profile, date=today).exists():
                _log_qr_scan_attempt(
                    request,
                    session=session,
                    user=request.user,
                    role=role,
                    profile=profile,
                    slot_no=slot_no,
                    accepted=False,
                    reason="Already marked for today",
                )
                return Response({"message": "Already Marked"}, status=200)

            mark_status = _current_mark_status()
            with transaction.atomic():
                TeacherAttendance.objects.create(teacher=profile, date=today, status=mark_status)
                _log_qr_scan_attempt(
                    request,
                    session=session,
                    user=request.user,
                    role=role,
                    profile=profile,
                    slot_no=slot_no,
                    accepted=True,
                    reason="Marked via QR",
                    marked_status=mark_status,
                )
            return Response({"message": f"Marked {mark_status}", "status": mark_status}, status=201)

        is_holiday = Holiday.objects.filter(date=today).filter(
            Q(applicable_for='everyone') | Q(applicable_for='staff')
        ).exists()
        if is_holiday:
            _log_qr_scan_attempt(
                request,
                session=session,
                user=request.user,
                role=role,
                profile=profile,
                slot_no=slot_no,
                accepted=False,
                reason="Holiday",
            )
            return Response({"error": "Today is a holiday. Attendance not allowed."}, status=400)

        if StaffAttendance.objects.filter(staff=profile, date=today).exists():
            _log_qr_scan_attempt(
                request,
                session=session,
                user=request.user,
                role=role,
                profile=profile,
                slot_no=slot_no,
                accepted=False,
                reason="Already marked for today",
            )
            return Response({"message": "Already Marked"}, status=200)

        mark_status = _current_mark_status()
        with transaction.atomic():
            StaffAttendance.objects.create(staff=profile, date=today, status=mark_status)
            _log_qr_scan_attempt(
                request,
                session=session,
                user=request.user,
                role=role,
                profile=profile,
                slot_no=slot_no,
                accepted=True,
                reason="Marked via QR",
                marked_status=mark_status,
            )
        return Response({"message": f"Marked {mark_status}", "status": mark_status}, status=201)
